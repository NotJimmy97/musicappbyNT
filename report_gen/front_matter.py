# -*- coding: utf-8 -*-
"""
front_matter.py
Sinh nội dung các trang đầu của Báo cáo Đồ án:
- Tiêu đề mở đầu (Ghi chú bìa)
- Nhận xét của Giảng viên hướng dẫn / Giảng viên chấm thi
- Lời cam đoan
- Lời cảm ơn
- Mục lục tổng hợp
- Danh mục từ viết tắt
- Danh mục bảng biểu
- Danh mục hình vẽ
"""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

def render_front_matter(builder, md_lines):
    # Tiêu đề báo cáo mở đầu
    builder.add_title(
        "BÁO CÁO ĐỒ ÁN MÔN HỌC: LẬP TRÌNH ỨNG DỤNG .NET",
        "Đề tài: XÂY DỰNG ỨNG DỤNG PHÁT NHẠC DESKTOP NATIVE TRÊN NỀN TẢNG WPF (.NET FRAMEWORK 4.6.1) VỚI KIẾN TRÚC MVVM, LOCAL BFF, AUDIO ENGINE DSP VÀ PERSISTENCE SQLITE"
    )

    md_lines.append("# BÁO CÁO ĐỒ ÁN MÔN HỌC: LẬP TRÌNH ỨNG DỤNG .NET\n")
    md_lines.append("## Đề tài: XÂY DỰNG ỨNG DỤNG PHÁT NHẠC DESKTOP NATIVE TRÊN NỀN TẢNG WPF (.NET FRAMEWORK 4.6.1) VỚI KIẾN TRÚC MVVM, LOCAL BFF, AUDIO ENGINE DSP VÀ PERSISTENCE SQLITE\n")
    md_lines.append("> *Ghi chú: Trang bìa chính và bìa phụ theo mẫu của Trường/Khoa sẽ được đính kèm tại thời điểm đóng tập chính thức.*\n")

    # =========================================================================
    # NHẬN XÉT CỦA GIẢNG VIÊN
    # =========================================================================
    builder.add_heading_1("NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN / CHẤM THI", page_break=True)
    md_lines.append("\n---\n\n# NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN / CHẤM THI\n")

    builder.add_paragraph("1. Về tiến độ thực hiện và tinh thần thái độ:")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("2. Về kiến trúc hệ thống và giải pháp kỹ thuật:")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("3. Về tính hoàn thiện của sản phẩm phần mềm:")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("4. Về hình thức trình bày và báo cáo:")
    builder.add_paragraph("................................................................................................................................................................")
    builder.add_paragraph("................................................................................................................................................................")
    
    # Bảng đánh giá điểm số
    builder.add_table(
        "Bảng 0.1: Phiếu đánh giá kết quả đồ án",
        ["Tiêu chí đánh giá", "Trọng số", "Điểm đánh giá (Thang 10)", "Chữ ký Giảng viên"],
        [
            ["Kiến trúc phần mềm (.NET / MVVM / BFF)", "25%", "", ""],
            ["Hiện thực kỹ thuật & Audio Engine DSP", "25%", "", ""],
            ["Thiết kế CSDL & Tối ưu SQLite", "20%", "", ""],
            ["Kiểm thử tự động (Unit Test / VSTest)", "15%", "", ""],
            ["Báo cáo thuyết minh & Thể thức văn bản", "15%", "", ""],
            ["TỔNG ĐIỂM CHUNG", "100%", "", ""]
        ]
    )

    md_lines.append("### Phiếu đánh giá kết quả đồ án\n")
    md_lines.append("| Tiêu chí đánh giá | Trọng số | Điểm đánh giá (Thang 10) | Chữ ký Giảng viên |\n|---|---|---|---|\n")
    md_lines.append("| Kiến trúc phần mềm (.NET / MVVM / BFF) | 25% | | |\n")
    md_lines.append("| Hiện thực kỹ thuật & Audio Engine DSP | 25% | | |\n")
    md_lines.append("| Thiết kế CSDL & Tối ưu SQLite | 20% | | |\n")
    md_lines.append("| Kiểm thử tự động (Unit Test / VSTest) | 15% | | |\n")
    md_lines.append("| Báo cáo thuyết minh & Thể thức văn bản | 15% | | |\n")
    md_lines.append("| TỔNG ĐIỂM CHUNG | 100% | | |\n\n")

    # =========================================================================
    # LỜI CAM ĐOAN
    # =========================================================================
    builder.add_heading_1("LỜI CAM ĐOAN", page_break=True)
    md_lines.append("\n---\n\n# LỜI CAM ĐOAN\n")

    p_camdoan = (
        "Em xin cam đoan đề tài đồ án môn học 'Xây dựng ứng dụng phát nhạc Desktop Native trên nền tảng WPF (.NET Framework 4.6.1)' "
        "là công trình nghiên cứu và thực nghiệm độc lập của chính bản thân em. Toàn bộ mã nguồn, cấu trúc giải thuật, kiến trúc phân tầng "
        "và các kết quả kiểm thử được trình bày trong báo cáo này đều phản ánh trung thực quá trình học tập và lập trình thực tế trên hệ thống. "
        "Các tài liệu tham khảo, các thư viện mã nguồn mở và nền tảng công nghệ được sử dụng (Microsoft .NET Framework, NAudio, SQLite, OWIN) "
        "đều đã được trích dẫn nguồn gốc và phiên bản rõ ràng theo đúng chuẩn mực học thuật."
    )
    builder.add_paragraph(p_camdoan)
    md_lines.append(p_camdoan + "\n\n")

    p_sign = builder.add_paragraph("Sinh viên thực hiện đề tài\n(Ký và ghi rõ họ tên)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    p_sign.runs[0].font.italic = True

    # =========================================================================
    # LỜI CẢM ƠN
    # =========================================================================
    builder.add_heading_1("LỜI CẢM ƠN", page_break=True)
    md_lines.append("\n---\n\n# LỜI CẢM ƠN\n")

    p_camon1 = (
        "Lời đầu tiên, em xin gửi lời cảm ơn chân thành và sâu sắc nhất đến Quý Thầy/Cô Bộ môn Công nghệ Phần mềm và Khoa Công nghệ Thông tin, "
        "những người đã tận tâm truyền đạt nền tảng kiến thức vững chắc về lập trình hướng đối tượng, tư duy kiến trúc hệ thống và các công nghệ cốt lõi "
        "của hệ sinh thái Microsoft .NET trong suốt học kỳ vừa qua."
    )
    p_camon2 = (
        "Đặc biệt, em xin bày tỏ lòng biết ơn sâu sắc đến Giảng viên hướng dẫn môn học, Thầy/Cô đã luôn định hướng phương pháp tiếp cận khoa học, "
        "đặt ra những yêu cầu khắt khe về tính chuẩn mực của kiến trúc phân lớp, nguyên lý tối ưu bộ nhớ zero-allocation và tác phong làm việc bài bản. "
        "Những lời góp ý quý báu của Thầy/Cô không chỉ giúp đồ án MusicApp đạt được độ hoàn thiện kỹ thuật cao mà còn rèn luyện cho em tư duy giải quyết "
        "vấn đề chuyên nghiệp của một kỹ sư phần mềm thực thụ."
    )
    builder.add_paragraph(p_camon1)
    builder.add_paragraph(p_camon2)
    md_lines.append(p_camon1 + "\n\n" + p_camon2 + "\n\n")

    # =========================================================================
    # MỤC LỤC TỔNG HỢP
    # =========================================================================
    builder.add_heading_1("MỤC LỤC TỔNG HỢP", page_break=True)
    md_lines.append("\n---\n\n# MỤC LỤC TỔNG HỢP\n")

    toc_items = [
        ("LỜI CAM ĐOAN", "i"),
        ("LỜI CẢM ƠN", "ii"),
        ("DANH MỤC TỪ VIẾT TẮT", "iv"),
        ("DANH MỤC BẢNG BIỂU", "v"),
        ("DANH MỤC HÌNH VẼ", "vi"),
        ("CHƯƠNG 1: TỔNG QUAN ĐỀ TÀI VÀ CÔNG NGHỆ ÁP DỤNG", "1"),
        ("    1.1. Đặt vấn đề và Mục tiêu đề tài", "1"),
        ("        1.1.1. Bối cảnh và Tính cấp thiết của đề tài", "1"),
        ("        1.1.2. Mục tiêu nghiên cứu và phát triển ứng dụng", "3"),
        ("        1.1.3. Phạm vi đề tài và giới hạn hệ thống", "4"),
        ("    1.2. Khảo sát nghiệp vụ và Yêu cầu hệ thống", "5"),
        ("        1.2.1. Khảo sát quy trình nghiệp vụ phát nhạc đa nguồn", "5"),
        ("        1.2.2. Phân tích yêu cầu chức năng (Functional Requirements)", "7"),
        ("        1.2.3. Phân tích yêu cầu phi chức năng (Non-Functional Requirements)", "9"),
        ("    1.3. Cơ sở công nghệ và Môi trường phát triển", "11"),
        ("        1.3.1. Nền tảng .NET Framework 4.6.1 và Ngôn ngữ C# 7.3", "11"),
        ("        1.3.2. Công nghệ giao diện WPF và Mô hình kiến trúc MVVM", "13"),
        ("        1.3.3. Mô hình Local Backend-for-Frontend (BFF) qua OWIN Self-Host", "15"),
        ("        1.3.4. Thư viện âm thanh NAudio và Xử lý tín hiệu số (DSP)", "17"),
        ("        1.3.5. Hệ quản trị CSDL SQLite và Mô hình Repository Pattern", "19"),
        ("        1.3.6. Môi trường phát triển và Công cụ hỗ trợ", "21"),
        ("CHƯƠNG 2: PHÂN TÍCH, THIẾT KẾ HỆ THỐNG VÀ CƠ SỞ DỮ LIỆU", "22"),
        ("    2.1. Phân tích chức năng và Thiết kế Use Case", "22"),
        ("        2.1.1. Sơ đồ Use Case tổng quát của hệ thống", "22"),
        ("        2.1.2. Đặc tả chi tiết các Use Case cốt lõi", "24"),
        ("        2.1.3. Thiết kế Sơ đồ tuần tự (Sequence Diagram)", "28"),
        ("        2.1.4. Thiết kế Sơ đồ hoạt động (Activity Diagram)", "31"),
        ("    2.2. Thiết kế Cơ sở dữ liệu (Database Design)", "33"),
        ("        2.2.1. Mô hình thực thể kết hợp (ERD - Entity Relationship Diagram)", "33"),
        ("        2.2.2. Sơ đồ quan hệ dữ liệu vật lý (Physical Schema Diagram)", "35"),
        ("        2.2.3. Bảng từ điển dữ liệu chi tiết (Data Dictionary)", "37"),
        ("    2.3. Thiết kế Kiến trúc phần mềm (.NET Solution Structure)", "41"),
        ("        2.3.1. Cấu trúc phân tầng dự án (Layered Solution Architecture)", "41"),
        ("        2.3.2. Sơ đồ chuỗi xử lý tín hiệu âm thanh DSP Pipeline", "44"),
        ("        2.3.3. Sơ đồ lớp chi tiết (Class Diagram)", "46"),
        ("CHƯƠNG 3: CÀI ĐẶT THỰC NGHIỆM VÀ KẾT QUẢ ĐẠT ĐƯỢC", "50"),
        ("    3.1. Môi trường triển khai và Cấu hình hệ thống", "50"),
        ("        3.1.1. Yêu cầu cấu hình triển khai", "50"),
        ("        3.1.2. Cấu hình chuỗi kết nối và Cơ chế Database Initializer", "51"),
        ("    3.2. Hiện thực hóa các Chức năng chính và Giao diện (Demo)", "53"),
        ("        3.2.1. Chức năng Tìm kiếm trực tuyến và Kỹ thuật Debounce 300ms", "53"),
        ("        3.2.2. Chức năng Phát nhạc, Quản lý Hàng đợi và Thẻ Now Playing", "55"),
        ("        3.2.3. Chức năng Quét thư viện cục bộ bằng giải thuật BFS", "58"),
        ("        3.2.4. Chức năng Đồng bộ Lời bài hát Karaoke (LrcParser)", "60"),
        ("        3.2.5. Chức năng Bộ cân bằng âm thanh DSP Equalizer 10 băng tần", "62"),
        ("        3.2.6. Chức năng Quản lý Playlist và Hệ thống Gợi ý bài hát", "64"),
        ("    3.3. Kiểm thử phần mềm (Testing)", "67"),
        ("        3.3.1. Chiến lược và Kế hoạch kiểm thử tự động", "67"),
        ("        3.3.2. Bảng kịch bản kiểm thử tiêu biểu (MSTest v2)", "68"),
        ("        3.3.3. Đánh giá kết quả kiểm thử toàn diện", "72"),
        ("    3.4. Đánh giá và Hướng phát triển", "73"),
        ("        3.4.1. Đánh giá ưu điểm nổi bật của hệ thống", "73"),
        ("        3.4.2. Những điểm hạn chế kỹ thuật còn tồn tại", "75"),
        ("        3.4.3. Đề xuất hướng nâng cấp và phát triển mở rộng", "76"),
        ("KẾT LUẬN", "78"),
        ("TÀI LIỆU THAM KHẢO", "80"),
        ("PHỤ LỤC", "82")
    ]

    for title, page in toc_items:
        p_toc = builder.doc.add_paragraph()
        p_toc.paragraph_format.line_spacing = 1.2
        p_toc.paragraph_format.space_before = Pt(1)
        p_toc.paragraph_format.space_after = Pt(2)
        r_title = p_toc.add_run(title)
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(12)
        if title.startswith("CHƯƠNG") or title in ["LỜI CAM ĐOAN", "LỜI CẢM ƠN", "KẾT LUẬN", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC"]:
            r_title.font.bold = True
        
        md_lines.append(f"- {title.strip()}\n")

    # =========================================================================
    # DANH MỤC TỪ VIẾT TẮT
    # =========================================================================
    builder.add_heading_1("DANH MỤC TỪ VIẾT TẮT", page_break=True)
    md_lines.append("\n---\n\n# DANH MỤC TỪ VIẾT TẮT\n")

    abbr_data = [
        ["WPF", "Windows Presentation Foundation", "Nền tảng giao diện đồ họa bản địa trên hệ điều hành Windows dựa trên DirectX."],
        ["XAML", "Extensible Application Markup Language", "Ngôn ngữ đánh dấu dựa trên XML dùng để khai báo giao diện người dùng."],
        ["MVVM", "Model - View - ViewModel", "Mô hình kiến trúc phần mềm tách biệt giữa Giao diện, Trạng thái và Dữ liệu."],
        ["BFF", "Backend-For-Frontend", "Mô hình kiến trúc gateway cục bộ trung gian phục vụ riêng cho ứng dụng client."],
        ["OWIN", "Open Web Interface for .NET", "Đặc tả chuẩn mở cho phép máy chủ web và ứng dụng .NET phân tách độc lập."],
        ["DSP", "Digital Signal Processing", "Xử lý tín hiệu kỹ thuật số (bộ lọc cân bằng âm sắc, phân tích âm học)."],
        ["EQ", "Equalizer", "Bộ cân bằng âm thanh điều chỉnh biên độ các dải tần số khác nhau."],
        ["FFT", "Fast Fourier Transform", "Thuật toán biến đổi Fourier nhanh chuyển đổi tín hiệu từ miền thời gian sang miền tần số."],
        ["WASAPI", "Windows Audio Session API", "Giao diện lập trình âm thanh mức thấp độ trễ cực nhỏ của hệ điều hành Windows."],
        ["CAS", "Content-Addressable Storage", "Cơ chế lưu trữ định danh dữ liệu trên đĩa dựa trên mã băm nội dung."],
        ["LRU", "Least Recently Used", "Thuật toán giải phóng bộ nhớ đệm dựa trên việc loại bỏ phần tử ít dùng gần đây nhất."],
        ["ADO.NET", "ActiveX Data Objects for .NET", "Công nghệ truy xuất dữ liệu mức thấp hướng kết nối và hiệu năng cao trong .NET."],
        ["DTO", "Data Transfer Object", "Đối tượng truyền tải dữ liệu giữa các tầng kiến trúc không chứa logic nghiệp vụ."],
        ["POCO", "Plain Old CLR Object", "Lớp đối tượng thuần túy trong C# không phụ thuộc vào bất kỳ framework cụ thể nào."],
        ["BFS", "Breadth-First Search", "Giải thuật duyệt theo chiều rộng áp dụng cho cây thư mục hệ thống tập tin."],
        ["ID3", "Identify an MP3", "Chuẩn thẻ siêu dữ liệu nhúng bên trong tệp âm thanh (Tên bài hát, Nghệ sĩ, Album, Ảnh bìa)."],
        ["WAL", "Write-Ahead Logging", "Cơ chế ghi log nhật ký trước giúp SQLite cho phép Đọc và Ghi đồng thời không khóa."],
        ["SOLID", "Single, Open-closed, Liskov, Interface, Dependency", "Bộ 5 nguyên lý thiết kế hướng đối tượng kinh điển trong kỹ nghệ phần mềm."]
    ]

    builder.add_table(
        "Bảng 0.2: Danh mục các từ viết tắt chuyên ngành trong báo cáo",
        ["Từ viết tắt", "Cụm từ tiếng Anh đầy đủ", "Ý nghĩa và vai trò kỹ thuật trong đề tài"],
        abbr_data
    )

    md_lines.append("| Từ viết tắt | Cụm từ tiếng Anh đầy đủ | Ý nghĩa và vai trò kỹ thuật trong đề tài |\n|---|---|---|\n")
    for abbr, full, desc in abbr_data:
        md_lines.append(f"| **{abbr}** | {full} | {desc} |\n")
    md_lines.append("\n")

    # =========================================================================
    # DANH MỤC BẢNG BIỂU
    # =========================================================================
    builder.add_heading_1("DANH MỤC BẢNG BIỂU", page_break=True)
    md_lines.append("\n---\n\n# DANH MỤC BẢNG BIỂU\n")

    tables_list = [
        ("Bảng 0.1", "Phiếu đánh giá kết quả đồ án"),
        ("Bảng 0.2", "Danh mục các từ viết tắt chuyên ngành trong báo cáo"),
        ("Bảng 1.1", "Bảng phân tích yêu cầu chức năng hệ thống (Functional Requirements)"),
        ("Bảng 1.2", "Bảng phân tích yêu cầu phi chức năng hệ thống (Non-Functional Requirements)"),
        ("Bảng 1.3", "So sánh đặc tính kỹ thuật: ADO.NET thuần vs. Entity Framework"),
        ("Bảng 1.4", "Thông số kỹ thuật bộ lọc cân bằng âm sắc 10 băng tần ISO"),
        ("Bảng 2.1", "Từ điển dữ liệu bảng tracks (Danh mục bài hát hợp nhất)"),
        ("Bảng 2.2", "Từ điển dữ liệu bảng library_folders và stream_cache"),
        ("Bảng 2.3", "Từ điển dữ liệu bảng playlists và playlist_tracks"),
        ("Bảng 2.4", "Từ điển dữ liệu bảng play_queue và user_interactions"),
        ("Bảng 2.5", "Từ điển dữ liệu bảng eq_presets và app_settings"),
        ("Bảng 2.6", "Phân bổ trách nhiệm các Project thành phần trong Solution"),
        ("Bảng 3.1", "Bảng 10 Kịch bản kiểm thử tiêu biểu trích xuất từ 60 bài test MSTest v2"),
        ("Bảng 3.2", "Ma trận so sánh tính năng MusicApp với các ứng dụng nghe nhạc hiện hành")
    ]

    for tid, tname in tables_list:
        p_t = builder.doc.add_paragraph()
        p_t.paragraph_format.line_spacing = 1.2
        r1 = p_t.add_run(f"{tid}: ")
        r1.font.bold = True
        p_t.add_run(tname)
        md_lines.append(f"- **{tid}**: {tname}\n")

    # =========================================================================
    # DANH MỤC HÌNH VẼ
    # =========================================================================
    builder.add_heading_1("DANH MỤC HÌNH VẼ", page_break=True)
    md_lines.append("\n---\n\n# DANH MỤC HÌNH VẼ\n")

    figures_list = [
        ("Hình 1.1", "Sơ đồ kiến trúc tổng thể 4 tầng của hệ thống MusicApp Desktop"),
        ("Hình 1.2", "Mô hình luồng dữ liệu MVVM và tương tác Dispatcher trong WPF"),
        ("Hình 2.1", "Sơ đồ Use Case tổng quát toàn bộ hệ thống MusicApp"),
        ("Hình 2.2", "Sơ đồ tuần tự (Sequence Diagram) - Luồng tìm kiếm bài hát trực tuyến"),
        ("Hình 2.3", "Sơ đồ tuần tự (Sequence Diagram) - Luồng phát nhạc và ghi bộ nhớ đệm CAS"),
        ("Hình 2.4", "Sơ đồ hoạt động (Activity Diagram) - Luồng quét thư viện cục bộ BFS"),
        ("Hình 2.5", "Sơ đồ thực thể liên kết (ERD) Cơ sở dữ liệu SQLite"),
        ("Hình 2.6", "Sơ đồ vật lý Cơ sở dữ liệu (Physical Schema Diagram) với các chỉ mục Index"),
        ("Hình 2.7", "Sơ đồ đường ống xử lý tín hiệu âm thanh kỹ thuật số (DSP Pipeline) trong NAudio"),
        ("Hình 2.8", "Sơ đồ lớp (Class Diagram) - Hệ thống ViewModel và Data Binding"),
        ("Hình 2.9", "Sơ đồ lớp (Class Diagram) - Hệ thống Persistence và Repository Pattern"),
        ("Hình 3.1", "Giao diện màn hình Tìm kiếm và Khám phá bài hát trực tuyến"),
        ("Hình 3.2", "Giao diện thẻ Now Playing và Trực quan hóa phổ âm thanh FFT 16 cột"),
        ("Hình 3.3", "Giao diện Hàng đợi phát nhạc (Play Queue) và cơ chế sắp xếp"),
        ("Hình 3.4", "Giao diện Quét và Quản lý thư viện âm nhạc cục bộ"),
        ("Hình 3.5", "Giao diện Hiển thị và Đồng bộ lời bài hát Karaoke (LRC Sync)"),
        ("Hình 3.6", "Giao diện Bộ cân bằng âm sắc DSP Equalizer 10 băng tần"),
        ("Hình 3.7", "Giao diện Quản lý Danh sách phát (Playlists) và Gợi ý thông minh"),
        ("Hình 3.8", "Biểu đồ kết quả thực thi 60 bài kiểm thử đơn vị MSTest v2 (100% Pass)")
    ]

    for fid, fname in figures_list:
        p_f = builder.doc.add_paragraph()
        p_f.paragraph_format.line_spacing = 1.2
        r1 = p_f.add_run(f"{fid}: ")
        r1.font.bold = True
        p_f.add_run(fname)
        md_lines.append(f"- **{fid}**: {fname}\n")
